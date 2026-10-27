from __future__ import annotations

import argparse
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


WHOLE_DOLLAR = Decimal("1")
DAYS_PER_YEAR = Decimal("365")
INTEREST_USER = "Interest"
REQUIRED_COLUMNS = ("date", "transaction", "account_balance", "user")


@dataclass(frozen=True)
class LedgerRow:
    occurred_on: date
    amount: Decimal
    user: str


@dataclass(frozen=True)
class Account:
    name: str
    opened_on: date
    opening_balance: Decimal
    opening_user: str
    rows: tuple[LedgerRow, ...]


@dataclass(frozen=True)
class MonthlyRow:
    year: int
    month: int
    account_balance: Decimal
    total_interest: Decimal
    user_balances: dict[str, Decimal]
    user_interests: dict[str, Decimal]

    @property
    def label(self) -> str:
        return f"{self.year}/{self.month}"


@dataclass(frozen=True)
class InterestCharge:
    """A bank interest charge compared with the interest calculated for it."""

    charged_on: date
    year: int
    month: int
    charged: Decimal
    calculated: Decimal


@dataclass(frozen=True)
class AccountReport:
    name: str
    annual_rate: Decimal
    users: tuple[str, ...]
    months: tuple[MonthlyRow, ...]
    charges: tuple[InterestCharge, ...]


def statement_window(year: int, month: int) -> tuple[date, date]:
    """Return the 21st-to-20th statement period payable in year/month."""
    if not 1 <= month <= 12:
        raise ValueError("month must be between 1 and 12")
    previous_year = year if month > 1 else year - 1
    previous_month = month - 1 if month > 1 else 12
    return date(previous_year, previous_month, 21), date(year, month, 20)


def statement_month(day: date) -> tuple[int, int]:
    """Return the (year, month) whose statement period contains ``day``."""
    if day.day >= 21:
        return (day.year + 1, 1) if day.month == 12 else (day.year, day.month + 1)
    return day.year, day.month


def previous_month(year: int, month: int) -> tuple[int, int]:
    return (year - 1, 12) if month == 1 else (year, month - 1)


def next_month(year: int, month: int) -> tuple[int, int]:
    return (year + 1, 1) if month == 12 else (year, month + 1)


def round_preserving_sum(
    values: dict[str, Decimal], target: Decimal
) -> dict[str, Decimal]:
    """Round each value to whole dollars so that the results add up to ``target``.

    Values are rounded half-up first; any leftover dollars go to the values
    that lost the most (or gained the least) in rounding.
    """
    target = target.quantize(WHOLE_DOLLAR, rounding=ROUND_HALF_UP)
    rounded = {
        key: value.quantize(WHOLE_DOLLAR, rounding=ROUND_HALF_UP)
        for key, value in values.items()
    }
    difference = target - sum(rounded.values(), Decimal("0"))
    if not rounded:
        if difference:
            raise ValueError(f"cannot split {target} across no users")
        return rounded
    step = WHOLE_DOLLAR if difference > 0 else -WHOLE_DOLLAR
    order = sorted(
        rounded,
        key=lambda key: values[key] - rounded[key],
        reverse=difference > 0,
    )
    for index in range(int(abs(difference))):
        rounded[order[index % len(order)]] += step
    return rounded


def calculate_account(
    account: Account, annual_rate: Decimal | str | float, as_of: date
) -> AccountReport:
    """Build the monthly table for every statement month that ended before ``as_of``."""
    try:
        rate = Decimal(str(annual_rate))
    except InvalidOperation as exc:
        raise ValueError("annual rate must be a number") from exc
    if not rate.is_finite() or rate < 0:
        raise ValueError("annual rate must be a finite, non-negative decimal")
    daily_rate = rate / DAYS_PER_YEAR

    users = _borrowers(account)
    rows_by_day: dict[date, list[LedgerRow]] = {}
    for row in account.rows:
        rows_by_day.setdefault(row.occurred_on, []).append(row)

    first_month = statement_month(account.opened_on)
    last_month = statement_month(as_of - timedelta(days=1))
    if statement_window(*last_month)[1] >= as_of:
        last_month = previous_month(*last_month)

    balances = {user: Decimal("0") for user in users}
    balances[account.opening_user] += account.opening_balance
    account_balance = account.opening_balance
    # Unrounded interest per statement month: the account total and each user's share.
    month_totals: dict[tuple[int, int], Decimal] = {}
    month_shares: dict[tuple[int, int], dict[str, Decimal]] = {}
    months: list[MonthlyRow] = []
    charges: list[InterestCharge] = []

    current_day = account.opened_on
    # Walk past the last finished month too, so later bank charges get checked.
    end_day = max([statement_window(*last_month)[1], *rows_by_day])
    while current_day <= end_day:
        current_month = statement_month(current_day)
        for row in rows_by_day.get(current_day, ()):
            if row.user == INTEREST_USER:
                charge = _split_interest_charge(
                    account.name, row, month_totals, month_shares, users
                )
                charges.append(charge.record)
                for user, share in charge.shares.items():
                    balances[user] += share
            else:
                balances[row.user] += row.amount
            account_balance += row.amount

        owed = -account_balance
        totals = month_shares.setdefault(
            current_month, {user: Decimal("0") for user in users}
        )
        month_totals.setdefault(current_month, Decimal("0"))
        if owed > 0:
            daily_interest = owed * daily_rate
            month_totals[current_month] += daily_interest
            debts = {user: -balance for user, balance in balances.items() if balance < 0}
            total_debt = sum(debts.values(), Decimal("0"))
            for user, debt in debts.items():
                totals[user] += daily_interest * debt / total_debt

        if (
            current_day == statement_window(*current_month)[1]
            and current_month <= last_month
        ):
            months.append(
                _monthly_row(
                    current_month,
                    account_balance,
                    balances,
                    month_totals[current_month],
                    totals,
                )
            )
        current_day += timedelta(days=1)

    return AccountReport(
        name=account.name,
        annual_rate=rate,
        users=users,
        months=tuple(months),
        charges=tuple(charges),
    )


@dataclass(frozen=True)
class _ChargeSplit:
    record: InterestCharge
    shares: dict[str, Decimal]


def _split_interest_charge(
    account_name: str,
    row: LedgerRow,
    month_totals: dict[tuple[int, int], Decimal],
    month_shares: dict[tuple[int, int], dict[str, Decimal]],
    users: tuple[str, ...],
) -> _ChargeSplit:
    """Split a bank charge by each user's interest in the month it pays for."""
    # The charge pays for the latest statement month that ended before it.
    billed_month = previous_month(*statement_month(row.occurred_on))
    shares = {
        user: share
        for user, share in month_shares.get(billed_month, {}).items()
        if share
    }
    total = sum(shares.values(), Decimal("0"))
    if not total:
        raise ValueError(
            f"sheet '{account_name}': interest charge on "
            f"{row.occurred_on.isoformat()} has no calculated interest "
            f"in {billed_month[0]}/{billed_month[1]} to split"
        )
    split = round_preserving_sum(
        {user: row.amount * share / total for user, share in shares.items()},
        row.amount,
    )
    return _ChargeSplit(
        record=InterestCharge(
            charged_on=row.occurred_on,
            year=billed_month[0],
            month=billed_month[1],
            charged=-row.amount,
            calculated=month_totals[billed_month].quantize(
                WHOLE_DOLLAR, rounding=ROUND_HALF_UP
            ),
        ),
        shares={user: split.get(user, Decimal("0")) for user in users},
    )


def _monthly_row(
    month: tuple[int, int],
    account_balance: Decimal,
    balances: dict[str, Decimal],
    interest: Decimal,
    shares: dict[str, Decimal],
) -> MonthlyRow:
    # Interest is a cost, so it is shown negative like the debt it comes from.
    total_interest = -interest.quantize(WHOLE_DOLLAR, rounding=ROUND_HALF_UP)
    return MonthlyRow(
        year=month[0],
        month=month[1],
        account_balance=account_balance.quantize(WHOLE_DOLLAR, rounding=ROUND_HALF_UP),
        total_interest=total_interest,
        user_balances=round_preserving_sum(balances, account_balance),
        user_interests=round_preserving_sum(
            {user: -share for user, share in shares.items()}, total_interest
        ),
    )


def _borrowers(account: Account) -> tuple[str, ...]:
    users = [account.opening_user]
    for row in account.rows:
        if row.user != INTEREST_USER and row.user not in users:
            users.append(row.user)
    return tuple(users)


def load_loan_workbook(path: str | Path, sheet: str | None = None) -> dict[str, Account]:
    """Read every sheet (or only ``sheet``) as one bank account."""
    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise ValueError(f"Excel file does not exist: {workbook_path}")

    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        return {
            worksheet.title: _read_sheet(worksheet, workbook.epoch)
            for worksheet in _select_worksheets(workbook, sheet)
        }
    finally:
        workbook.close()


def _select_worksheets(workbook: object, sheet: str | None) -> list[object]:
    if sheet is None:
        return list(workbook.worksheets)
    if sheet not in workbook.sheetnames:
        raise ValueError(f"sheet does not exist: {sheet}")
    return [workbook[sheet]]


def _read_sheet(worksheet: object, epoch: datetime) -> Account:
    title = worksheet.title
    headers = _header_columns(worksheet)
    missing = [name for name in REQUIRED_COLUMNS if name not in headers]
    if missing:
        raise ValueError(f"sheet '{title}' is missing columns: {', '.join(missing)}")

    opening: tuple[date, Decimal, str] | None = None
    rows: list[LedgerRow] = []
    running_balance = Decimal("0")
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=2, values_only=True), start=2
    ):
        raw = {name: _cell(row, headers[name]) for name in REQUIRED_COLUMNS}
        if all(value is None for value in raw.values()):
            continue
        if raw["date"] is None:
            raise ValueError(f"sheet '{title}', row {row_number}: Date is required")
        occurred_on = _parse_excel_date(raw["date"], epoch, title, row_number)
        user = _parse_user(raw["user"], title, row_number)
        recorded_balance = (
            None
            if raw["account_balance"] is None
            else _parse_excel_amount(raw["account_balance"], "Account_balance", title, row_number)
        )

        if raw["transaction"] is None:
            # Only the first row may omit Transaction: it states the opening balance.
            if opening is not None:
                raise ValueError(
                    f"sheet '{title}', row {row_number}: Transaction is required"
                )
            if recorded_balance is None or user == INTEREST_USER:
                raise ValueError(
                    f"sheet '{title}', row {row_number}: the opening row needs an "
                    "Account_balance and a borrower in User"
                )
            opening = (occurred_on, recorded_balance, user)
            running_balance = recorded_balance
            continue

        amount = _parse_excel_amount(raw["transaction"], "Transaction", title, row_number)
        if opening is None:
            opening = (occurred_on, Decimal("0"), user if user != INTEREST_USER else title)
        if rows and occurred_on < rows[-1].occurred_on:
            raise ValueError(
                f"sheet '{title}', row {row_number}: rows must be in date order"
            )
        running_balance += amount
        if recorded_balance is not None and recorded_balance != running_balance:
            raise ValueError(
                f"sheet '{title}', row {row_number}: Account_balance {recorded_balance} "
                f"does not match the running total {running_balance}"
            )
        rows.append(LedgerRow(occurred_on=occurred_on, amount=amount, user=user))

    if opening is None:
        raise ValueError(f"sheet '{title}' contains no transactions")
    opened_on, opening_balance, opening_user = opening
    return Account(
        name=title,
        opened_on=opened_on,
        opening_balance=opening_balance,
        opening_user=opening_user,
        rows=tuple(rows),
    )


def _cell(row: tuple, index: int | None) -> object:
    if index is None or index >= len(row):
        return None
    return row[index]


def _parse_user(raw: object, sheet: str, row: int) -> str:
    user = "" if raw is None else str(raw).strip()
    if not user:
        raise ValueError(f"sheet '{sheet}', row {row}: User is required")
    return user


def _header_columns(worksheet: object) -> dict[str, int]:
    first_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    return {
        str(value).strip().lower(): index
        for index, value in enumerate(first_row)
        if value is not None
    }


def _parse_excel_date(raw: object, epoch: datetime, sheet: str, row: int) -> date:
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    if isinstance(raw, (int, float, Decimal)) and not isinstance(raw, bool):
        converted = from_excel(raw, epoch)
        return converted.date() if isinstance(converted, datetime) else converted
    if isinstance(raw, str):
        text = raw.strip()
        for date_format in ("%Y/%m/%d", "%Y-%m-%d"):
            try:
                return datetime.strptime(text, date_format).date()
            except ValueError:
                pass
    raise ValueError(f"sheet '{sheet}', row {row}: invalid Date value {raw!r}")


def _parse_excel_amount(raw: object, column: str, sheet: str, row: int) -> Decimal:
    if isinstance(raw, bool):
        raise ValueError(f"sheet '{sheet}', row {row}: invalid {column} value {raw!r}")
    try:
        amount = Decimal(str(raw).replace(",", "").strip())
    except (InvalidOperation, AttributeError) as exc:
        raise ValueError(
            f"sheet '{sheet}', row {row}: invalid {column} value {raw!r}"
        ) from exc
    if not amount.is_finite():
        raise ValueError(f"sheet '{sheet}', row {row}: {column} must be finite")
    return amount


def calculate_reports(
    accounts: dict[str, Account],
    rates: dict[str | None, Decimal],
    as_of: date,
) -> list[AccountReport]:
    """Calculate every account; ``rates[None]`` is the rate for unnamed sheets."""
    reports = []
    for name, account in accounts.items():
        rate = rates.get(name, rates.get(None))
        if rate is None:
            raise ValueError(f"no rate given for sheet '{name}' (use --rate {name}=RATE)")
        reports.append(calculate_account(account, rate, as_of))
    return reports


def rate_argument(value: str) -> tuple[str | None, Decimal]:
    """Parse ``RATE`` (every sheet) or ``SHEET=RATE`` (one sheet)."""
    sheet, separator, text = value.rpartition("=")
    try:
        rate = Decimal(text.strip())
    except InvalidOperation as exc:
        raise argparse.ArgumentTypeError("rate must be a number") from exc
    if not rate.is_finite() or rate < 0:
        raise argparse.ArgumentTypeError("rate must be a non-negative decimal")
    if separator and not sheet.strip():
        raise argparse.ArgumentTypeError("sheet name before '=' is empty")
    return (sheet.strip() if separator else None), rate


def format_money(value: Decimal) -> str:
    rounded = value.quantize(WHOLE_DOLLAR, rounding=ROUND_HALF_UP)
    # Adding zero turns a rounded "-0" into "0".
    return f"{rounded + 0:,.0f}"


def _display_width(text: str) -> int:
    # CJK characters take two terminal columns.
    return sum(2 if unicodedata.east_asian_width(char) in "WF" else 1 for char in text)


def _pad(text: str, width: int, align_right: bool) -> str:
    padding = " " * (width - _display_width(text))
    return padding + text if align_right else text + padding


def format_table(report: AccountReport) -> str:
    headers = [
        "Month",
        "Account_balance",
        "Total_interest",
        *(f"{user}_balance" for user in report.users),
        *(f"{user}_interest" for user in report.users),
    ]
    body = [
        [
            row.label,
            format_money(row.account_balance),
            format_money(row.total_interest),
            *(format_money(row.user_balances[user]) for user in report.users),
            *(format_money(row.user_interests[user]) for user in report.users),
        ]
        for row in report.months
    ]
    widths = [
        max(_display_width(line[index]) for line in [headers, *body])
        for index in range(len(headers))
    ]

    def render(cells: list[str], header: bool) -> str:
        return "".join(
            f" {_pad(cell, width, align_right=index > 0 and not header)} |"
            for index, (cell, width) in enumerate(zip(cells, widths))
        )

    separator = "".join(f"-{'-' * width}-+" for width in widths)
    return "\n".join([render(headers, True), separator, *(render(line, False) for line in body)])


def format_report(report: AccountReport) -> str:
    lines = [f"帳戶: {report.name}（年利率 {(report.annual_rate * 100).normalize():f}%）"]
    if report.months:
        lines.append(format_table(report))
    else:
        lines.append("尚無已結束的計息月份。")
    for charge in report.charges:
        if charge.charged != charge.calculated:
            lines.append(
                f"注意: {charge.charged_on.isoformat()} 銀行扣息 "
                f"{format_money(charge.charged)}，與 {charge.year}/{charge.month} "
                f"計算利息 {format_money(charge.calculated)} 不同；"
                "使用者餘額依實際扣息分攤。"
            )
    return "\n".join(lines)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "依 Excel 帳戶紀錄，計算每月（上月 21 日至本月 20 日）銀行利息，"
            "並依各使用者的借款拆帳。"
        )
    )
    parser.add_argument(
        "--input", default="loan.xlsx", help="Excel 路徑（預設：loan.xlsx）"
    )
    parser.add_argument(
        "--rate",
        required=True,
        action="append",
        type=rate_argument,
        help=(
            "年利率（小數格式），例如 0.0292 代表 2.92%%；"
            "用 SHEET=RATE 指定單一帳戶，例如 --rate CK2=0.0292 --rate CK3=0.0275"
        ),
    )
    parser.add_argument(
        "--sheet", help="只讀取指定的 sheet（例如 CK3）；預設讀取全部 sheet"
    )
    return parser


def main(argv: list[str] | None = None, today: date | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        accounts = load_loan_workbook(args.input, args.sheet)
        reports = calculate_reports(accounts, dict(args.rate), today or date.today())
    except ValueError as exc:
        parser.error(str(exc))

    print("\n\n".join(format_report(report) for report in reports))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
