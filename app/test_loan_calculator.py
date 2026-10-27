from __future__ import annotations

import argparse
import tempfile
import unittest
from contextlib import redirect_stdout
from datetime import date
from decimal import Decimal
from io import StringIO
from pathlib import Path

from openpyxl import Workbook

from loan_calculator import (
    Account,
    LedgerRow,
    calculate_account,
    load_loan_workbook,
    main,
    rate_argument,
    round_preserving_sum,
    statement_window,
)


HEADER = ["Date", "Transaction", "Account_balance", "User", "Note"]
AS_OF = date(2026, 12, 1)


def account(*rows: tuple[date, int, str], opened_on: date | None = None) -> Account:
    ledger = tuple(LedgerRow(day, Decimal(amount), user) for day, amount, user in rows)
    return Account(
        name="CK",
        opened_on=opened_on or ledger[0].occurred_on,
        opening_balance=Decimal("0"),
        opening_user=ledger[0].user,
        rows=ledger,
    )


def month(report, year: int, month_number: int):
    return next(row for row in report.months if (row.year, row.month) == (year, month_number))


class HelperTests(unittest.TestCase):
    def test_statement_window_crosses_year(self) -> None:
        self.assertEqual(
            statement_window(2026, 1),
            (date(2025, 12, 21), date(2026, 1, 20)),
        )

    def test_rate_argument_accepts_global_and_per_sheet_rates(self) -> None:
        self.assertEqual(rate_argument("0.0292"), (None, Decimal("0.0292")))
        self.assertEqual(rate_argument("CK3=0.0275"), ("CK3", Decimal("0.0275")))
        for invalid in ("x", "-0.1", "=0.1"):
            with self.subTest(invalid=invalid):
                with self.assertRaises(argparse.ArgumentTypeError):
                    rate_argument(invalid)

    def test_round_preserving_sum_hits_the_target(self) -> None:
        values = {"A": Decimal("1.4"), "B": Decimal("1.4"), "C": Decimal("1.2")}
        rounded = round_preserving_sum(values, Decimal("4"))
        self.assertEqual(sum(rounded.values()), Decimal("4"))
        # The leftover dollar goes to the value that lost the most in rounding.
        self.assertEqual(rounded, {"A": Decimal("2"), "B": Decimal("1"), "C": Decimal("1")})


class AccountInterestTests(unittest.TestCase):
    def test_daily_interest_uses_closing_balance_and_statement_window(self) -> None:
        # 100,000 at 2.92% costs 8 a day. 5/10-5/11 at 100,000, 5/12-5/20 at 50,000.
        report = calculate_account(
            account(
                (date(2026, 5, 10), -100000, "Allen"),
                (date(2026, 5, 12), 50000, "Allen"),
            ),
            "0.0292",
            AS_OF,
        )
        may = month(report, 2026, 5)
        self.assertEqual(may.total_interest, Decimal("-52"))
        self.assertEqual(may.account_balance, Decimal("-50000"))
        self.assertEqual(may.user_interests["Allen"], Decimal("-52"))

    def test_same_day_borrow_and_repay_costs_nothing(self) -> None:
        report = calculate_account(
            account(
                (date(2026, 7, 12), -100000, "Allen"),
                (date(2026, 7, 12), 100000, "Allen"),
            ),
            "0.0292",
            AS_OF,
        )
        self.assertEqual(month(report, 2026, 7).total_interest, Decimal("0"))

    def test_loan_on_21st_is_billed_next_month(self) -> None:
        report = calculate_account(
            account((date(2026, 4, 21), -36500, "Allen"), opened_on=date(2026, 4, 1)),
            "0.03",
            AS_OF,
        )
        self.assertEqual(month(report, 2026, 4).total_interest, Decimal("0"))
        self.assertEqual(month(report, 2026, 5).total_interest, Decimal("-90"))

    def test_interest_is_split_by_each_users_daily_debt(self) -> None:
        # 7/1-7/10: Allen owes 100,000 alone (80). 7/11-7/20: both owe 100,000 (80 each).
        report = calculate_account(
            account(
                (date(2026, 7, 1), -100000, "Allen"),
                (date(2026, 7, 11), -100000, "Ray"),
            ),
            "0.0292",
            AS_OF,
        )
        july = month(report, 2026, 7)
        self.assertEqual(july.total_interest, Decimal("-240"))
        self.assertEqual(july.user_interests, {"Allen": Decimal("-160"), "Ray": Decimal("-80")})
        self.assertEqual(july.user_balances, {"Allen": Decimal("-100000"), "Ray": Decimal("-100000")})

    def test_bank_charge_is_split_like_the_interest_it_pays(self) -> None:
        report = calculate_account(
            account(
                (date(2026, 7, 1), -100000, "Allen"),
                (date(2026, 7, 11), -100000, "Ray"),
                (date(2026, 7, 21), -240, "Interest"),
            ),
            "0.0292",
            AS_OF,
        )
        august = month(report, 2026, 8)
        self.assertEqual(august.account_balance, Decimal("-200240"))
        self.assertEqual(
            august.user_balances,
            {"Allen": Decimal("-100160"), "Ray": Decimal("-100080")},
        )
        self.assertNotIn("Interest", report.users)
        self.assertEqual(report.charges[0].charged, report.charges[0].calculated)

    def test_users_always_add_up_to_account_totals(self) -> None:
        report = calculate_account(
            account(
                (date(2026, 6, 25), -33333, "A"),
                (date(2026, 6, 26), -33333, "B"),
                (date(2026, 6, 27), -33334, "C"),
                (date(2026, 7, 21), -64, "Interest"),
                (date(2026, 7, 30), 10000, "B"),
            ),
            "0.0292",
            AS_OF,
        )
        for row in report.months:
            with self.subTest(month=row.label):
                self.assertEqual(sum(row.user_balances.values()), row.account_balance)
                self.assertEqual(sum(row.user_interests.values()), row.total_interest)

    def test_unfinished_month_is_not_reported(self) -> None:
        report = calculate_account(
            account((date(2026, 7, 1), -100000, "Allen"), opened_on=date(2026, 6, 1)),
            "0.0292",
            date(2026, 7, 20),
        )
        self.assertEqual([row.label for row in report.months], ["2026/6"])

    def test_interest_charge_without_calculated_interest_is_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "no calculated interest"):
            calculate_account(
                account(
                    (date(2026, 7, 1), 0, "Allen"),
                    (date(2026, 7, 21), -10, "Interest"),
                ),
                "0.0292",
                AS_OF,
            )


class WorkbookTests(unittest.TestCase):
    def make_workbook(self, path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "CK3"
        sheet.append(HEADER)
        sheet.append([date(2026, 7, 1), None, 0, "Olivia", None])
        sheet.append([date(2026, 7, 1), -100000, -100000, "Allen", None])
        sheet.append([date(2026, 7, 11), -100000, -200000, "Ray", None])
        sheet.append([date(2026, 7, 21), -240, -200240, "Interest", None])
        other = workbook.create_sheet("CK2")
        other.append(HEADER)
        other.append([date(2026, 7, 1), -50000, -50000, "Allen", "note is ignored"])
        workbook.save(path)

    def test_each_sheet_is_one_account(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "loan.xlsx"
            self.make_workbook(path)
            accounts = load_loan_workbook(path)
        self.assertEqual(list(accounts), ["CK3", "CK2"])
        self.assertEqual(accounts["CK3"].opening_user, "Olivia")
        self.assertEqual(len(accounts["CK3"].rows), 3)
        self.assertEqual(accounts["CK2"].rows[0].amount, Decimal("-50000"))

    def test_account_balance_must_match_running_total(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "loan.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(HEADER)
            sheet.append([date(2026, 7, 1), -100, -99, "Allen", None])
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "does not match"):
                load_loan_workbook(path)

    def test_unknown_sheet_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "loan.xlsx"
            self.make_workbook(path)
            with self.assertRaisesRegex(ValueError, "sheet does not exist"):
                load_loan_workbook(path, "Nope")

    def test_cli_prints_one_aligned_table_per_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "loan.xlsx"
            self.make_workbook(path)
            output = StringIO()
            with redirect_stdout(output):
                return_code = main(
                    ["--input", str(path), "--rate", "0.0292", "--sheet", "CK3"],
                    today=date(2026, 7, 25),
                )
        self.assertEqual(return_code, 0)
        lines = output.getvalue().splitlines()
        self.assertIn("帳戶: CK3", lines[0])
        self.assertEqual(
            lines[1].split("|")[:-1],
            [" Month  ", " Account_balance ", " Total_interest ", " Olivia_balance ",
             " Allen_balance ", " Ray_balance ", " Olivia_interest ", " Allen_interest ",
             " Ray_interest "],
        )
        self.assertEqual(len({len(line) for line in lines[1:]}), 1)
        self.assertIn("-240", lines[-1])
        self.assertNotIn("CK2", output.getvalue())

    def test_cli_requires_a_rate_for_every_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "loan.xlsx"
            self.make_workbook(path)
            with redirect_stdout(StringIO()), self.assertRaises(SystemExit):
                main(["--input", str(path), "--rate", "CK3=0.0275"], today=AS_OF)


if __name__ == "__main__":
    unittest.main()
